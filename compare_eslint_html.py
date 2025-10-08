# -*- coding: utf-8 -*-
import os
import re
import csv
from bs4 import BeautifulSoup
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
import sys


def safe_remove(path):
    try:
        if os.path.exists(path):
            os.remove(path)
    except PermissionError:
        print(f" Excel 文件 {path} 正在被占用，请关闭后再运行。")


def canonicalize_path(path: str) -> str:
    """标准化路径用于比较"""
    return path.replace('/', '\\').strip().lower()


def should_include(path: str) -> bool:
    """忽略 build-work 文件"""
    return 'build-work' not in path.lower()


def clean_file(th):
    raw = ''.join(th.find_all(string=True, recursive=False)).strip()
    raw = re.sub(r'^[\[\]\+\-\s]+', '', raw)
    return raw.strip()


def parse_summary(text):
    m = re.search(r'(\d+)\s*problems?\s*\(\s*(\d+)\s*errors?,\s*(\d+)\s*warnings?\s*\)', text, re.I)
    return tuple(map(int, m.groups())) if m else (None, None, None)


def parse_html(path):
    """解析 ESLint HTML 报告"""
    with open(path, 'r', encoding='utf-8') as f:
        soup = BeautifulSoup(f, 'lxml')

    result = {}
    for tr in soup.select('tr[data-group]'):
        group = tr.get('data-group')
        th = tr.find('th')
        if not th:
            continue

        file_path = clean_file(th)
        if not should_include(file_path):
            continue

        norm_path = canonicalize_path(file_path)
        span = th.find('span')
        summary_text = span.get_text(strip=True) if span else ''
        summary_tuple = parse_summary(summary_text)

        rows = []
        sib = tr.find_next_sibling('tr')
        while sib and not sib.has_attr('data-group'):
            if group in sib.get('class', []):
                tds = sib.find_all('td')
                if tds:
                    line_col = tds[0].get_text(strip=True)
                    line_num = line_col.split(':')[0] if ':' in line_col else line_col
                    sev = tds[1].get_text(strip=True)
                    msg = tds[2].get_text(" ", strip=True)
                    rule = tds[3].get_text(strip=True) if len(tds) > 3 else ''
                    rows.append((line_num, sev, msg, rule))
            sib = sib.find_next_sibling('tr')

        result[norm_path] = {
            'display_path': file_path,
            'summary_tuple': summary_tuple,
            'summary_text': summary_text,
            'rows': rows
        }
    return result


def compare_reports(old_path, new_path,
                    output_csv='eslint_new_issues.csv',
                    output_xlsx='eslint_new_issues.xlsx'):
    old = parse_html(old_path)
    new = parse_html(new_path)
    diff_rows = []

    for key in sorted(new.keys()):
        n = new[key]
        o = old.get(key, {'summary_tuple': (0, 0, 0), 'summary_text': 'N/A', 'rows': []})

        n_sum, o_sum = n['summary_tuple'], o['summary_tuple']

        # 只当 summary 有增加时才继续
        def greater(a, b):
            def val(x): return x if isinstance(x, int) else -1
            return any(val(a[i]) > val(b[i]) for i in range(3))

        if not greater(n_sum, o_sum):
            continue

        # 保留所有重复行（不使用 set）
        new_rows = n['rows']
        old_rows = o['rows']

        for r in new_rows:
            if r not in old_rows:
                diff_rows.append([
                    n['display_path'],
                    o['summary_text'],
                    n['summary_text'],
                    r[0], r[1], r[2], r[3]
                ])

    headers = ['file', 'summary_old', 'summary_new', 'line', 'severity', 'message', 'rule']

    safe_remove(output_xlsx)

    # 写入 CSV
    with open(output_csv, 'w', newline='', encoding='utf-8-sig') as f:
        csv.writer(f).writerows([headers] + diff_rows)

    if not diff_rows:
        print(" 没有新增问题（或数字未增加）。")
        return

    df = pd.DataFrame(diff_rows, columns=headers)
    with pd.ExcelWriter(output_xlsx, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='New Issues')

    wb = load_workbook(output_xlsx)
    ws = wb.active

    # 合并相同文件名（顶部对齐）
    merge_col = 1
    cur, start = None, 2
    for r in range(2, ws.max_row + 2):
        val = ws.cell(r, merge_col).value
        if val != cur:
            if start < r - 1:
                ws.merge_cells(start_row=start, start_column=merge_col,
                               end_row=r - 1, end_column=merge_col)
                ws.cell(start, merge_col).alignment = Alignment(vertical='top', horizontal='left')
            cur, start = val, r

    # 自动列宽
    for col in ws.columns:
        max_len = max((len(str(c.value)) if c.value else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 120)

    wb.save(output_xlsx)
    print(f" 新增问题已写入：{output_csv} 和 {output_xlsx}（共 {len(df)} 条）")


if __name__ == '__main__':
    # 获取 exe 运行路径（支持打包与源码两种模式）
    if getattr(sys, 'frozen', False):
        base_dir = os.path.dirname(sys.executable)
    else:
        base_dir = os.path.dirname(os.path.abspath(__file__))

    old_dir = os.path.join(base_dir, "old")
    new_dir = os.path.join(base_dir, "new")
    out_dir = os.path.join(base_dir, "output")
    os.makedirs(out_dir, exist_ok=True)

    # 若目录不存在则创建
    for folder in [old_dir, new_dir]:
        os.makedirs(folder, exist_ok=True)

    old_files = [f for f in os.listdir(old_dir) if f.endswith('.html')]
    new_files = [f for f in os.listdir(new_dir) if f.endswith('.html')]

    if not old_files or not new_files:
        print(" 请在 old/ 和 new/ 文件夹中放入 .html 报告文件后重试。")
        input("按 Enter 键退出...")
    else:
        old_path = os.path.join(old_dir, old_files[0])
        new_path = os.path.join(new_dir, new_files[0])
        print(f" Old: {old_path}")
        print(f" New: {new_path}")

        output_csv = os.path.join(out_dir, "eslint_new_issues.csv")
        output_xlsx = os.path.join(out_dir, "eslint_new_issues.xlsx")

        compare_reports(old_path, new_path, output_csv, output_xlsx)
        input("\n 任务完成，按 Enter 键退出...")
